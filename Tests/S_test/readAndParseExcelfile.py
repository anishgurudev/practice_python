#Method 1: openpyxl — Cell-by-Cell Control
# Logic: Load the workbook
# →select a sheet
# →iterate rows and access each cell by .value.

from openpyxl import load_workbook

# Step 1: Load the workbook
wb = load_workbook("testdata.xlsx")
print(wb)
# Step 2: Select a sheet
ws = wb.active
print(ws)

# print("Sheet name:", ws.title)
# print("Total rows:", ws.max_row)
# print("Total cols:", ws.max_column)